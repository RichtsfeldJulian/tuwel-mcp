"""Convert downloaded course files and Moodle HTML into plain text for agents."""

from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup

TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".py", ".java", ".c", ".h",
    ".cpp", ".hpp", ".js", ".ts", ".sql", ".r", ".m", ".tex", ".sh", ".ipynb", ".rs", ".go",
    ".hs", ".kt", ".cs", ".php", ".rb", ".scala", ".pl", ".lisp", ".asm", ".vhd", ".log",
}


def html_to_text(html: str | None) -> str:
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    for a in soup.find_all("a", href=True):
        text = a.get_text(strip=True)
        if text and a["href"] not in text:
            a.replace_with(f"{text} ({a['href']})")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for block in soup.find_all(["p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"]):
        block.insert_after("\n")
    text = soup.get_text()
    text = re.sub(r"[ \t\xa0]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_text(path: Path, max_pages: int | None = None) -> str | None:
    """Return extracted text, or None if the format is not supported."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = reader.pages if max_pages is None else reader.pages[:max_pages]
        parts = []
        for i, page in enumerate(pages, 1):
            try:
                parts.append(f"--- page {i} ---\n{page.extract_text() or ''}")
            except Exception as exc:  # malformed pages should not kill the whole document
                parts.append(f"--- page {i} --- [extraction failed: {exc}]")
        if max_pages is not None and len(reader.pages) > max_pages:
            parts.append(f"[... {len(reader.pages) - max_pages} more pages not shown]")
        return "\n".join(parts)
    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            parts.append(f"--- slide {i} ---\n" + "\n".join(texts))
        return "\n".join(parts)
    if suffix in (".html", ".htm"):
        return html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    return None
